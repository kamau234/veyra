"""Generation of the authoritative VEYRA product import workbook.

The template produced here is the single definition of the import contract:
``TEMPLATE_COLUMNS`` is consumed by the reader and the validator so column
names, order and dropdown choices can never drift apart.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.constants import UNITS
from core.exceptions import FileError
from core.logger import get_logger

logger = get_logger("imports.template")

#: Every column the template writes, in order. One definition — the reader and
#: the validator import this so names and order can never drift apart.
TEMPLATE_COLUMNS = (
    "Product Code",
    "Product Name",
    "Category",
    "Subcategory",
    "Brand",
    "Unit",
    "Cost Price",
    "Selling Price",
    "Opening Stock",
    "Reorder Level",
    "VAT Applicable",
)

#: Columns a workbook must carry to be read at all (blueprint 7.2).
REQUIRED_COLUMNS = (
    "Product Code",
    "Product Name",
    "Category",
    "Unit",
    "Cost Price",
    "Selling Price",
    "Opening Stock",
    "Reorder Level",
    "VAT Applicable",
)

#: Optional hierarchy levels. A workbook without these headers (one saved from
#: an older template) still imports; the levels are simply left blank.
OPTIONAL_COLUMNS = ("Subcategory", "Brand")

TEMPLATE_TITLE = "VEYRA PRODUCT IMPORT TEMPLATE"
BRAND_FILL = "2563EB"
HEADER_FILL = "DBEAFE"
BORDER_COLOR = "9CA3AF"

#: Rows the dropdowns cover; users rarely paste more than a few hundred products.
DROPDOWN_LAST_ROW = 500

_EXAMPLE_ROWS = (
    ("COF001", "Classic", "Beverages", "Coffee", "Nescafé", "Piece", 480, 570, 24, 6, "Yes"),
    ("COF002", "Gold", "Beverages", "Coffee", "Nescafé", "Piece", 850, 990, 12, 4, "Yes"),
    ("COF003", "House Blend", "Beverages", "Coffee", "Dormans", "Pack", 620, 750, 18, 5, "Yes"),
    ("DET001", "1kg Washing Powder", "Detergents", "Washing Powder", "Ariel", "Piece", 290, 350, 40, 10, "Yes"),
    ("DET002", "1kg Washing Powder", "Detergents", "Washing Powder", "Omo", "Piece", 260, 320, 35, 10, "Yes"),
    ("DIS001", "Dishwashing Liquid 750ml", "Detergents", "Dishwashing", "Sunlight", "Bottle", 210, 265, 30, 8, "Yes"),
)

_COLUMN_WIDTHS = {
    "Product Code": 16,
    "Product Name": 32,
    "Category": 18,
    "Subcategory": 20,
    "Brand": 16,
    "Unit": 12,
    "Cost Price": 14,
    "Selling Price": 14,
    "Opening Stock": 15,
    "Reorder Level": 15,
    "VAT Applicable": 15,
}

INSTRUCTIONS = (
    "1. Do not rename the required columns.",
    "2. Enter one product per row.",
    "3. Category is the broad group the product belongs to, e.g. Beverages or Detergents.",
    "4. Subcategory is the specific group inside that category, e.g. Coffee or Washing Powder.",
    "5. Brand is the manufacturer or brand, e.g. Nescafé, Dormans, Ariel or Sunlight.",
    "6. Product Name is the specific product or line, e.g. Classic 100g or House Blend 250g.",
    "7. Product Code is the unique identifier (SKU) for that product; no two products share one.",
    "8. Subcategory and Brand may be left blank when they genuinely do not apply.",
    "9. Never put the brand in the Category column, and do not repeat the brand in the "
    "Product Name when the Brand column already carries it.",
    "10. Category, Subcategory and Brand can be new or existing values; matching names are "
    "grouped together automatically.",
    "11. Prices must be numeric; do not type currency symbols into cells.",
    "12. Opening Stock and Reorder Level must be non-negative numbers.",
    "13. VAT Applicable must be Yes or No.",
    "14. Save the workbook as .xlsx before importing.",
)

#: Column name / required / example, printed on the Instructions sheet.
_COLUMN_REFERENCE = (
    ("Product Code", "Yes", "COF001"),
    ("Product Name", "Yes", "Classic"),
    ("Category", "Yes", "Beverages"),
    ("Subcategory", "No", "Coffee"),
    ("Brand", "No", "Nescafé"),
    ("Unit", "Yes", "Piece"),
    ("Cost Price", "Yes", "480"),
    ("Selling Price", "Yes", "570"),
    ("Opening Stock", "Yes", "24"),
    ("Reorder Level", "Yes", "6"),
    ("VAT Applicable", "Yes", "Yes"),
)

#: What each level means, printed above the column reference on Instructions.
HIERARCHY_GUIDE = (
    ("Level", "What it holds", "Example"),
    ("Category", "The broad group a product belongs to.", "Beverages"),
    ("Subcategory", "The specific group inside that category.", "Coffee"),
    ("Brand", "The manufacturer or brand of the product.", "Nescafé"),
    ("Product Name", "The specific product or product line.", "Classic"),
    ("Product Code", "The unique identifier (SKU) for that product.", "COF001"),
)


def _build_products_sheet(sheet) -> None:
    last_column = len(TEMPLATE_COLUMNS)
    banner_font = Font(bold=True, size=14, color="FFFFFF")
    banner_fill = PatternFill("solid", fgColor=BRAND_FILL)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    for column in range(1, last_column + 1):
        cell = sheet.cell(row=1, column=column)
        cell.fill = banner_fill
        cell.font = banner_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.cell(row=1, column=1, value=TEMPLATE_TITLE)
    sheet.row_dimensions[1].height = 26

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, name in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = sheet.cell(row=2, column=column, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = _COLUMN_WIDTHS[name]

    for offset, example in enumerate(_EXAMPLE_ROWS):
        for column, value in enumerate(example, start=1):
            sheet.cell(row=3 + offset, column=column, value=value).border = border

    sheet.freeze_panes = "A3"

    unit_column = get_column_letter(TEMPLATE_COLUMNS.index("Unit") + 1)
    vat_column = get_column_letter(TEMPLATE_COLUMNS.index("VAT Applicable") + 1)
    unit_validation = DataValidation(
        type="list",
        formula1='"' + ",".join(UNITS) + '"',
        allow_blank=True,
        showErrorMessage=True,
    )
    unit_validation.error = "Choose a unit from the list."
    unit_validation.errorTitle = "Invalid unit"
    sheet.add_data_validation(unit_validation)
    unit_validation.add(f"{unit_column}3:{unit_column}{DROPDOWN_LAST_ROW}")

    vat_validation = DataValidation(
        type="list",
        formula1='"Yes,No"',
        allow_blank=True,
        showErrorMessage=True,
    )
    vat_validation.error = "VAT Applicable must be Yes or No."
    vat_validation.errorTitle = "Invalid value"
    sheet.add_data_validation(vat_validation)
    vat_validation.add(f"{vat_column}3:{vat_column}{DROPDOWN_LAST_ROW}")


def _build_instructions_sheet(sheet) -> None:
    title = sheet.cell(row=1, column=1, value="How to fill in the import template")
    title.font = Font(bold=True, size=13, color=BRAND_FILL)

    row = 3
    for line in INSTRUCTIONS:
        sheet.cell(row=row, column=1, value=line)
        row += 1

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)

    row += 1
    sheet.cell(row=row, column=1, value="How products are organised").font = Font(bold=True, size=12)
    row += 1
    for offset, line in enumerate(HIERARCHY_GUIDE):
        for column, value in enumerate(line, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            if offset == 0:
                cell.font = header_font
                cell.fill = header_fill
        row += 1

    row += 1
    for column, heading in enumerate(("Column", "Required", "Example"), start=1):
        cell = sheet.cell(row=row, column=column, value=heading)
        cell.font = header_font
        cell.fill = header_fill
    row += 1
    for reference in _COLUMN_REFERENCE:
        for column, value in enumerate(reference, start=1):
            sheet.cell(row=row, column=column, value=value)
        row += 1

    sheet.column_dimensions["A"].width = 60
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 28


def generate_template(destination: Path | str) -> Path:
    """Write the official import template to ``destination`` and return the path."""
    path = Path(destination)
    workbook = Workbook()
    products = workbook.active
    products.title = "Products"
    _build_products_sheet(products)
    _build_instructions_sheet(workbook.create_sheet("Instructions"))
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    except OSError as exc:
        logger.error("Could not save the import template to %s", path, exc_info=exc)
        raise FileError(
            f"The template could not be saved to '{path.name}'.", detail=repr(exc)
        ) from exc
    logger.info("Generated import template at %s", path)
    return path
