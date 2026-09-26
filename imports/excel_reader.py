"""Read a product workbook into raw rows.

This layer only understands spreadsheet structure — locating the header row
and pairing cells with canonical column names. Business rules (prices, units,
duplicates) belong to :mod:`imports.excel_validator`.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from core.exceptions import FileError
from core.logger import get_logger
from imports.excel_template import TEMPLATE_COLUMNS

logger = get_logger("imports.reader")

#: How many leading rows to scan for the header (the template puts a banner
#: above it, and users often paste in extra title rows).
HEADER_SCAN_LIMIT = 15

#: A row is only a header candidate if it contains both of these.
_HEADER_MARKERS = ("product code", "product name")

_PREFERRED_SHEET = "Products"


@dataclass(frozen=True)
class RawRow:
    """One spreadsheet row.

    ``row_number`` is the 1-based Excel row the user sees so error messages
    line up with the file they open. ``values`` is keyed by canonical column
    name and holds cells as-is; normalization is the validator's job.
    """

    row_number: int
    values: dict[str, object]


@dataclass(frozen=True)
class RawWorkbook:
    headers: list[str]
    rows: list[RawRow]
    source: Path
    sheet_name: str


def _normalized(value: object) -> str:
    return " ".join(str(value or "").split()).lower()


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    """Return (header index, canonical column -> column index) for the header row."""
    canonical_by_text = {_normalized(name): name for name in TEMPLATE_COLUMNS}
    for index, row in enumerate(rows[:HEADER_SCAN_LIMIT]):
        texts = {_normalized(value) for value in row}
        if not all(marker in texts for marker in _HEADER_MARKERS):
            continue
        column_map: dict[str, int] = {}
        for position, value in enumerate(row):
            name = canonical_by_text.get(_normalized(value))
            # Keep the first column when a name appears more than once.
            if name is not None and name not in column_map:
                column_map[name] = position
        return index, column_map
    raise FileError(
        "Could not find a header row containing 'Product Code' and 'Product Name'. "
        "Use the VEYRA import template and do not rename the columns."
    )


def _is_blank(values: dict[str, object]) -> bool:
    return all(value is None or str(value).strip() == "" for value in values.values())


def read_workbook(path: Path | str) -> RawWorkbook:
    """Parse an .xlsx file into raw rows without applying business rules."""
    source = Path(path)
    if not source.exists():
        raise FileError(f"'{source.name}' was not found.")
    if source.suffix.lower() != ".xlsx":
        raise FileError(
            f"'{source.name}' is not an .xlsx file. Save the workbook as .xlsx before importing."
        )

    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = (
            workbook[_PREFERRED_SHEET]
            if _PREFERRED_SHEET in workbook.sheetnames
            else workbook.active
        )
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        sheet_name = sheet.title
        workbook.close()
    except FileError:
        raise
    except Exception as exc:  # noqa: BLE001 - corrupt/foreign files surface as FileError
        logger.error("Could not parse workbook %s", source, exc_info=exc)
        raise FileError(
            f"'{source.name}' could not be opened as an Excel workbook.",
            detail=repr(exc),
        ) from exc

    header_index, column_map = _find_header(rows)
    headers = [name for name in TEMPLATE_COLUMNS if name in column_map]

    raw_rows: list[RawRow] = []
    for index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = {
            name: (row[position] if position < len(row) else None)
            for name, position in column_map.items()
        }
        if _is_blank(values):
            continue
        raw_rows.append(RawRow(row_number=index, values=values))

    return RawWorkbook(
        headers=headers, rows=raw_rows, source=source, sheet_name=sheet_name
    )
