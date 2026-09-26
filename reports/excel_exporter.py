"""XLSX export for reports.

Numbers stay numbers (money as ``#,##0.00``, counts as ``#,##0``) so the owner
can total or pivot an exported sheet in Excel without re-typing anything.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.dates import format_datetime
from core.exceptions import FileError
from core.logger import get_logger
from core.money import to_money
from services.reporting_service import ReportColumn, ReportResult

logger = get_logger("reports.xlsx")

CURRENCY_LINE = "Currency: KES (KSh)"
EMPTY_MESSAGE = "No records for the selected period."

HEADER_FILL = "2563EB"
BORDER_COLOR = "E5E7EB"
TOTALS_FILL = "F3F4F6"

_NUMBER_FORMATS = {"money": "#,##0.00", "int": "#,##0"}


def _sheet_title(report: ReportResult) -> str:
    # Excel caps sheet names at 31 characters and forbids : \ / ? * [ ]
    cleaned = "".join(char for char in report.title if char not in ':\\/?*[]')
    return (cleaned or "Report")[:31]


def _cell_value(column: ReportColumn, value) -> object:
    """Native Excel value for a cell, so formatting and maths still work."""
    if value is None or value == "":
        return None
    if column.kind == "money":
        return float(to_money(value))
    if column.kind == "int":
        return int(value)
    if column.kind in ("datetime", "date"):
        return value if isinstance(value, (datetime, date)) else str(value)
    return str(value)


def _width(column: ReportColumn, rows: list[dict]) -> int:
    widest = max(
        [len(column.label)]
        + [len(column.format(row.get(column.key))) for row in rows[:200]]
    )
    return min(max(widest + 3, 10), 42)


def export_xlsx(
    report: ReportResult,
    destination: Path | str,
    *,
    business_name: str,
    logo_path: str | None = None,
) -> Path:
    """Write ``report`` to ``destination`` as a formatted workbook."""
    path = Path(destination)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(report)

    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    last_column = max(len(report.columns), 1)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = sheet.cell(row=1, column=1, value=business_name)
    title_cell.font = Font(bold=True, size=14, color=HEADER_FILL)

    meta = [report.title]
    if report.subtitle:
        meta.append(report.subtitle)
    if report.period_label:
        meta.append(f"Period: {report.period_label}")
    meta.append(f"Generated: {format_datetime(report.generated_at)}")
    meta.append(CURRENCY_LINE)

    row_index = 2
    for line in meta:
        sheet.cell(row=row_index, column=1, value=line).font = Font(
            bold=(row_index == 2), color="6B7280" if row_index > 2 else "111827"
        )
        row_index += 1
    row_index += 1

    header_row = row_index
    for position, column in enumerate(report.columns, start=1):
        cell = sheet.cell(row=header_row, column=position, value=column.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.border = border
        cell.alignment = Alignment(
            horizontal="right" if column.align == "right" else "left", vertical="center"
        )
        sheet.column_dimensions[get_column_letter(position)].width = _width(
            column, report.rows
        )

    data_row = header_row
    if report.is_empty:
        data_row += 1
        sheet.cell(row=data_row, column=1, value=EMPTY_MESSAGE).font = Font(
            italic=True, color="6B7280"
        )
    for row in report.rows:
        data_row += 1
        for position, column in enumerate(report.columns, start=1):
            cell = sheet.cell(
                row=data_row, column=position, value=_cell_value(column, row.get(column.key))
            )
            cell.border = border
            number_format = _NUMBER_FORMATS.get(column.kind)
            if number_format:
                cell.number_format = number_format
            elif column.kind in ("datetime", "date"):
                cell.number_format = (
                    "dd mmm yyyy hh:mm" if column.kind == "datetime" else "dd mmm yyyy"
                )
            cell.alignment = Alignment(horizontal=column.align)

    if report.totals and report.rows:
        data_row += 1
        for position, column in enumerate(report.columns, start=1):
            value = None if position == 1 else (report.totals or {}).get(column.key)
            cell = sheet.cell(
                row=data_row,
                column=position,
                value="TOTAL" if position == 1 else _cell_value(column, value),
            )
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=TOTALS_FILL)
            cell.border = border
            number_format = _NUMBER_FORMATS.get(column.kind)
            if number_format:
                cell.number_format = number_format
            cell.alignment = Alignment(horizontal=column.align)

    if report.notes:
        data_row += 2
        for note in report.notes:
            sheet.cell(row=data_row, column=1, value=note).font = Font(
                italic=True, color="6B7280"
            )
            data_row += 1

    if report.rows:
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(last_column)}{header_row + len(report.rows)}"
        )

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    except OSError as exc:
        logger.error("Excel export to %s failed: %s", path, exc, exc_info=exc)
        raise FileError("The Excel file could not be written.", detail=repr(exc)) from exc

    logger.info(
        "Excel export '%s': %d rows -> %s", report.title, report.row_count, path
    )
    return path
