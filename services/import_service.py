"""Excel import: preview, atomic commit and error reporting.

Blueprint 7.4/7.5. The commit is all-or-nothing, and an update only ever
touches product master data — live stock and images are never overwritten by
an import, because stock changes belong in Inventory.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from core.constants import MovementType, Reason
from core.exceptions import FileError
from core.logger import get_logger
from db.session import session_scope
from imports.excel_reader import read_workbook
from imports.excel_validator import ImportPreview, ImportRow, validate_workbook
from models import Product
from repositories.category_repository import CategoryRepository
from repositories.product_repository import ProductRepository
from services.inventory_service import apply_movement

logger = get_logger("services.import")


@dataclass(frozen=True)
class ImportSummary:
    imported: int
    updated: int
    skipped: int
    errors: int
    created_categories: list[str] = field(default_factory=list)

    @property
    def message(self) -> str:
        parts = [f"Imported {self.imported} product{'s' if self.imported != 1 else ''}"]
        if self.updated:
            parts.append(f"updated {self.updated}")
        if self.skipped:
            parts.append(f"skipped {self.skipped}")
        if self.errors:
            parts.append(f"{self.errors} row{'s' if self.errors != 1 else ''} with errors")
        summary = ", ".join(parts) + "."
        if self.created_categories:
            summary += " Created categories: " + ", ".join(self.created_categories) + "."
        return summary


def preview_import(path: Path | str) -> ImportPreview:
    """Read and validate a workbook without writing anything."""
    try:
        return validate_workbook(read_workbook(path))
    except FileError:
        raise
    except Exception as exc:  # noqa: BLE001 - user sees a message, log keeps detail
        logger.error("Import preview failed for %s", path, exc_info=exc)
        raise FileError("That workbook could not be read.", detail=repr(exc)) from exc


def commit_import(preview: ImportPreview) -> ImportSummary:
    """Write a validated preview to the database in one transaction."""
    if not preview.add and not preview.update:
        return ImportSummary(
            imported=0,
            updated=0,
            skipped=0,
            errors=len(preview.rejected),
        )

    created_categories: list[str] = []
    imported = updated = skipped = 0

    with session_scope() as session:
        products = ProductRepository(session)
        categories = CategoryRepository(session)

        for row in preview.add:
            if products.code_exists(row.code):
                logger.warning("Skipped row %s: code %s appeared during import", row.row_number, row.code)
                skipped += 1
                continue
            category, was_created = categories.get_or_create(row.category)
            if was_created and category is not None:
                created_categories.append(category.name)

            product = Product(
                code=row.code,
                name=row.name,
                category_id=category.id if category else None,
                unit=row.unit,
                cost_price=row.cost_price,
                selling_price=row.selling_price,
                stock_quantity=0,
                reorder_level=row.reorder_level,
                vat_applicable=row.vat_applicable,
                is_active=True,
            )
            products.add(product)
            if row.opening_stock > 0:
                apply_movement(
                    session,
                    product=product,
                    movement_type=MovementType.STOCK_IN,
                    reason=Reason.INITIAL_STOCK,
                    quantity=row.opening_stock,
                    direction="in",
                    reference=None,
                    notes="Opening stock from Excel import.",
                )
            imported += 1

        for row in preview.update:
            product = session.get(Product, row.existing_product_id)
            if product is None:
                skipped += 1
                continue
            if products.code_exists(row.code, exclude_id=product.id):
                logger.warning(
                    "Skipped row %s: code %s now belongs to another product",
                    row.row_number,
                    row.code,
                )
                skipped += 1
                continue
            _apply_master_data(session, product, row)
            updated += 1

    summary = ImportSummary(
        imported=imported,
        updated=updated,
        skipped=skipped,
        errors=len(preview.rejected),
        created_categories=created_categories,
    )
    logger.info(
        "Import committed: %s imported, %s updated, %s skipped, %s errors",
        imported,
        updated,
        skipped,
        summary.errors,
    )
    return summary


def _apply_master_data(session, product: Product, row: ImportRow) -> None:
    """Update master-data fields only. Stock and image are deliberately untouched."""
    product.code = row.code
    product.name = row.name
    product.unit = row.unit
    product.cost_price = row.cost_price
    product.selling_price = row.selling_price
    product.reorder_level = row.reorder_level
    product.vat_applicable = row.vat_applicable
    category, _ = CategoryRepository(session).get_or_create(row.category)
    product.category_id = category.id if category else product.category_id
    session.flush()


def write_error_report(preview: ImportPreview, destination: Path | str) -> Path:
    """Write the rejected rows to XLSX so the owner can fix and retry."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path = Path(destination)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import Errors"

    headers = ["Row", "Column", "Problem", "Product Code", "Product Name"]
    sheet.append(headers)
    for index, _ in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="DC2626")

    for rejected in preview.rejected:
        code = rejected.values.get("Product Code") or ""
        name = rejected.values.get("Product Name") or ""
        for error in rejected.errors:
            sheet.append([error.row, error.column, error.message, str(code), str(name)])

    widths = (8, 18, 60, 18, 32)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    sheet.freeze_panes = "A2"

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    except OSError as exc:
        logger.error("Could not write the import error report to %s", path, exc_info=exc)
        raise FileError("The error report could not be saved.", detail=repr(exc)) from exc

    logger.info("Import error report written to %s (%d rows)", path, len(preview.rejected))
    return path
