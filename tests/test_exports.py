"""Report and receipt exports: CSV, XLSX, PDF reports and the sale receipt."""

from __future__ import annotations

import csv
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from core.exceptions import FileError
from imports.excel_template import TEMPLATE_COLUMNS
from reports.csv_exporter import export_csv
from reports.excel_exporter import export_xlsx
from reports.pdf_receipt import export_receipt_pdf, render_receipt_text
from reports.pdf_reports import export_pdf
from services import (
    import_service,
    inventory_service,
    pos_service,
    product_service,
    reporting_service,
    settings_service,
)

SALES = "Sales Report"
INVENTORY = "Inventory Report"

CATALOGUE = (
    ("DET001", "Ariel Detergent 1kg", "Detergents", "Piece", 290, 350, 40, 10, "Yes"),
    ("BRD001", "Supreme Bread 400g", "Bakery", "Piece", 45, 60, 25, 8, "No"),
    ("HM001", "Hammer", "Hardware", "Piece", 700, 950, 3, 2, "Yes"),
)


def write_workbook(path, rows) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(list(TEMPLATE_COLUMNS))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)


def sheet_name(report) -> str:
    """Mirror the exporter's Excel sheet-name rules."""
    cleaned = "".join(char for char in report.title if char not in ':\\/?*[]')
    return (cleaned or "Report")[:31]


@pytest.fixture
def shop(tmp_path):
    """Three products, a configured business profile and one completed sale."""
    write_workbook(tmp_path / "catalogue.xlsx", CATALOGUE)
    import_service.commit_import(import_service.preview_import(tmp_path / "catalogue.xlsx"))

    settings_service.update_business_profile(
        business_name="Mama Njeri Shop",
        owner_name="Njeri W.",
        phone="0712 345 678",
        address="Nakuru, Kenya",
    )

    products = {product.code: product for product in product_service.list_products()}
    cart = pos_service.Cart()
    cart.add(products["DET001"], 2)
    cart.add(products["HM001"], 1)
    return pos_service.complete_sale(cart)


def report(name: str):
    return reporting_service.run_report(name)


def header_row_of(sheet, first_label: str) -> int:
    return next(
        index for index in range(1, 15) if sheet.cell(row=index, column=1).value == first_label
    )


class TestCsv:
    def test_writes_identification_block_headers_and_totals(self, shop, tmp_path):
        result = report(SALES)
        path = export_csv(result, tmp_path / "sales.csv", business_name="Mama Njeri Shop")

        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))

        assert rows[0] == ["Mama Njeri Shop"]
        assert rows[1] == [result.title]
        assert rows[4][0].startswith("Currency: KES")

        header = [column.label for column in result.columns]
        assert header in rows
        assert rows[-1][0] == "TOTAL"
        assert rows[-1][header.index("Total")] == f"{result.totals['total']:,.2f}"

    def test_bom_marks_the_file_as_utf8_for_excel(self, shop, tmp_path):
        path = export_csv(report(INVENTORY), tmp_path / "i.csv", business_name="Shop")
        assert path.read_bytes()[:3] == b"\xef\xbb\xbf"

    def test_empty_report_says_so_instead_of_writing_nothing(self, tmp_path):
        result = reporting_service.sales_report(period="Today")
        assert result.is_empty
        path = export_csv(result, tmp_path / "empty.csv", business_name="Shop")
        assert "No records for the selected period." in path.read_text(encoding="utf-8-sig")

    def test_unwritable_destination_raises_file_error(self, shop, tmp_path):
        blocker = tmp_path / "blocker"
        blocker.write_text("not a folder", encoding="utf-8")
        with pytest.raises(FileError):
            export_csv(report(SALES), blocker / "sales.csv", business_name="Shop")


class TestExcel:
    def test_money_stays_numeric_and_formatted(self, shop, tmp_path):
        result = report(SALES)
        path = export_xlsx(result, tmp_path / "sales.xlsx", business_name="Mama Njeri Shop")
        sheet = load_workbook(path)[sheet_name(result)]

        first = header_row_of(sheet, "Date")
        labels = [sheet.cell(row=first, column=index).value for index in range(1, 10)]
        total_column = labels.index("Total") + 1
        cell = sheet.cell(row=first + 1, column=total_column)

        assert isinstance(cell.value, (int, float)) and not isinstance(cell.value, str)
        assert cell.number_format == "#,##0.00"
        assert cell.value == pytest.approx(float(result.totals["total"]))

    def test_sheet_carries_the_identification_block(self, shop, tmp_path):
        result = report(INVENTORY)
        path = export_xlsx(result, tmp_path / "inv.xlsx", business_name="Mama Njeri Shop")
        sheet = load_workbook(path)[sheet_name(result)]
        block = " ".join(str(sheet.cell(row=index, column=1).value) for index in range(1, 9))

        assert "Mama Njeri Shop" in block
        assert result.title in block
        assert "Currency: KES (KSh)" in block

    def test_stock_quantities_are_written_as_integers(self, shop, tmp_path):
        result = report(INVENTORY)
        path = export_xlsx(result, tmp_path / "inv.xlsx", business_name="Shop")
        sheet = load_workbook(path)[sheet_name(result)]

        first = header_row_of(sheet, "Code")
        labels = [sheet.cell(row=first, column=index).value for index in range(1, 12)]
        stock = sheet.cell(row=first + 1, column=labels.index("Stock") + 1)
        assert isinstance(stock.value, int)
        assert stock.number_format == "#,##0"

    def test_totals_row_is_written(self, shop, tmp_path):
        result = report(SALES)
        sheet = load_workbook(
            export_xlsx(result, tmp_path / "s.xlsx", business_name="Shop")
        )[sheet_name(result)]
        assert "TOTAL" in [sheet.cell(row=index, column=1).value for index in range(1, 20)]

    def test_every_report_exports(self, shop, tmp_path):
        for name in reporting_service.REPORT_NAMES:
            path = export_xlsx(report(name), tmp_path / f"{name}.xlsx", business_name="Shop")
            assert path.stat().st_size > 0


class TestPdf:
    def test_every_report_produces_a_pdf(self, shop, tmp_path):
        for name in reporting_service.REPORT_NAMES:
            path = export_pdf(
                report(name), tmp_path / f"{name}.pdf", business_name="Mama Njeri Shop"
            )
            assert path.read_bytes()[:4] == b"%PDF"
            assert path.stat().st_size > 1000

    def test_empty_report_still_produces_a_pdf(self, tmp_path):
        result = reporting_service.sales_report(period="Today")
        path = export_pdf(result, tmp_path / "empty.pdf", business_name="Shop")
        assert path.read_bytes()[:4] == b"%PDF"

    def test_a_long_report_spans_several_pages(self, shop, tmp_path):
        product = product_service.find_by_code("DET001")
        for index in range(60):
            inventory_service.record_movement(
                product_id=product.id,
                movement_type="Stock In",
                reason="Restock",
                quantity=1,
                reference=f"GRN-{index:03d}",
            )

        result = reporting_service.stock_movement_report(period="This Month")
        assert result.row_count >= 60
        path = export_pdf(result, tmp_path / "long.pdf", business_name="Shop")
        assert path.read_bytes()[:4] == b"%PDF"
        assert path.stat().st_size > 3000

    def test_special_characters_are_escaped_not_interpreted(self, tmp_path):
        product_service.create_product(
            product_service.ProductInput(
                code="AMP1",
                name="A & B <Wires> 2.5mm",
                category="Electrical & Misc",
                unit="Metre",
                cost_price=Decimal("10"),
                selling_price=Decimal("20"),
                stock_quantity=5,
                reorder_level=1,
                vat_applicable=True,
            )
        )
        result = reporting_service.inventory_report()
        assert export_pdf(
            result, tmp_path / "amp.pdf", business_name="A & B Ltd"
        ).read_bytes()[:4] == b"%PDF"
        assert export_xlsx(result, tmp_path / "amp.xlsx", business_name="A & B Ltd").exists()
        assert export_csv(result, tmp_path / "amp.csv", business_name="A & B Ltd").exists()


class TestReceipt:
    def test_text_receipt_shows_lines_and_totals(self, shop):
        sale = reporting_service.get_sale(shop.sale_id)
        text = render_receipt_text(sale, settings_service.get_settings())

        assert "Mama Njeri Shop" in text
        assert sale.invoice_number in text
        assert "Ariel Detergent 1kg" in text
        assert "700.00" in text  # 2 x 350.00
        assert "950.00" in text  # 1 x 950.00
        assert "1,650.00" in text  # subtotal
        assert "264.00" in text  # VAT at 16%
        assert "1,914.00" in text  # total
        assert "TOTAL" in text

    def test_pdf_receipt_is_written(self, shop, tmp_path):
        sale = reporting_service.get_sale(shop.sale_id)
        path = export_receipt_pdf(
            sale, tmp_path / "receipt.pdf", settings=settings_service.get_settings()
        )
        assert path.read_bytes()[:4] == b"%PDF"
        assert path.stat().st_size > 800

    def test_receipt_reflects_the_snapshot_not_current_master_data(self, shop, tmp_path):
        """Blueprint 13.4: reprinting an old receipt must not show new prices."""
        product = product_service.find_by_code("DET001")
        product_service.update_product(
            product.id,
            product_service.ProductInput(
                code=product.code,
                name="Ariel Detergent 3kg RENAMED",
                category="Detergents",
                unit="Piece",
                cost_price=Decimal("400"),
                selling_price=Decimal("999"),
                stock_quantity=0,
                reorder_level=10,
                vat_applicable=True,
            ),
        )

        sale = reporting_service.get_sale(shop.sale_id)
        text = render_receipt_text(sale, settings_service.get_settings())
        assert "Ariel Detergent 1kg" in text
        assert "RENAMED" not in text
        assert "999.00" not in text
        assert export_receipt_pdf(
            sale, tmp_path / "old.pdf", settings=settings_service.get_settings()
        ).exists()

    def test_discount_appears_on_the_receipt(self, shop, tmp_path):
        products = {product.code: product for product in product_service.list_products()}
        cart = pos_service.Cart()
        cart.add(products["BRD001"], 10)
        result = pos_service.complete_sale(cart, discount=Decimal("100"))

        sale = reporting_service.get_sale(result.sale_id)
        text = render_receipt_text(sale, settings_service.get_settings())
        assert "Discount" in text
        assert "100.00" in text
        assert export_receipt_pdf(
            sale, tmp_path / "d.pdf", settings=settings_service.get_settings()
        ).exists()
