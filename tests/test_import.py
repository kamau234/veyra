"""Excel template generation, import validation and the atomic commit."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.exceptions import FileError
from imports.excel_template import (
    OPTIONAL_COLUMNS,
    REQUIRED_COLUMNS,
    TEMPLATE_COLUMNS,
    generate_template,
)
from services import import_service, inventory_service, product_service


def write_workbook(path: Path, rows: list[tuple], *, header: tuple = TEMPLATE_COLUMNS) -> Path:
    """Build a plain workbook with a header row and the given data rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


GOOD_ROW = (
    "DET001", "1kg Washing Powder", "Detergents", "Washing Powder", "Ariel",
    "Piece", 290, 350, 40, 10, "Yes",
)

#: The nine-column layout older templates produced.
LEGACY_HEADER = tuple(column for column in TEMPLATE_COLUMNS if column not in OPTIONAL_COLUMNS)
LEGACY_ROW = ("DET001", "Ariel Detergent 1kg", "Detergents", "Piece", 290, 350, 40, 10, "Yes")


class TestTemplate:
    def test_generates_a_workbook_with_both_sheets(self, tmp_path):
        path = generate_template(tmp_path / "template.xlsx")
        assert path.exists() and path.stat().st_size > 0

        workbook = load_workbook(path)
        assert "Products" in workbook.sheetnames
        assert "Instructions" in workbook.sheetnames

    def test_products_sheet_matches_the_documented_contract(self, tmp_path):
        path = generate_template(tmp_path / "template.xlsx")
        sheet = load_workbook(path)["Products"]

        assert sheet.cell(row=1, column=1).value == "VEYRA PRODUCT IMPORT TEMPLATE"
        assert tuple(
            sheet.cell(row=2, column=index).value for index in range(1, len(TEMPLATE_COLUMNS) + 1)
        ) == TEMPLATE_COLUMNS
        assert TEMPLATE_COLUMNS == (
            "Product Code", "Product Name", "Category", "Subcategory", "Brand", "Unit",
            "Cost Price", "Selling Price", "Opening Stock", "Reorder Level", "VAT Applicable",
        )
        # Six example rows, each with a code and a Yes/No VAT flag in the last column.
        last = len(TEMPLATE_COLUMNS)
        for row in range(3, 9):
            assert sheet.cell(row=row, column=1).value
            assert sheet.cell(row=row, column=last).value in {"Yes", "No"}
        assert sheet.freeze_panes == "A3"

    def test_example_rows_show_the_hierarchy(self, tmp_path):
        sheet = load_workbook(generate_template(tmp_path / "template.xlsx"))["Products"]
        examples = [
            tuple(sheet.cell(row=row, column=column).value for column in range(1, 6))
            for row in range(3, 9)
        ]
        assert ("COF001", "Classic", "Beverages", "Coffee", "Nescafé") in examples
        assert ("COF003", "House Blend", "Beverages", "Coffee", "Dormans") in examples
        assert ("DET001", "1kg Washing Powder", "Detergents", "Washing Powder", "Ariel") in examples
        assert ("DET002", "1kg Washing Powder", "Detergents", "Washing Powder", "Omo") in examples
        assert ("DIS001", "Dishwashing Liquid 750ml", "Detergents", "Dishwashing", "Sunlight") in examples

    def test_instructions_sheet_lists_the_rules(self, tmp_path):
        sheet = load_workbook(generate_template(tmp_path / "t.xlsx"))["Instructions"]
        text = " ".join(
            str(value)
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if value is not None
        )
        assert "Do not rename the required columns" in text
        assert "VAT Applicable must be Yes or No" in text
        assert "Save the workbook as .xlsx" in text
        # Each level is explained, and blanks are allowed where they make sense.
        assert "broad group" in text
        assert "specific group inside that category" in text
        assert "manufacturer or brand" in text
        assert "unique identifier (SKU)" in text
        assert "may be left blank when they genuinely do not apply" in text
        assert "Never put the brand in the Category column" in text

    def test_generated_template_can_be_imported_as_is(self, tmp_path):
        path = generate_template(tmp_path / "template.xlsx")
        preview = import_service.preview_import(path)
        assert preview.is_valid
        assert len(preview.add) == 6
        summary = import_service.commit_import(preview)
        assert summary.imported == 6
        assert product_service.catalogue_stats()["products"] == 6
        assert product_service.brand_names() == ["Ariel", "Dormans", "Nescafé", "Omo", "Sunlight"]


class TestReader:
    def test_missing_file_raises_file_error(self, tmp_path):
        with pytest.raises(FileError):
            import_service.preview_import(tmp_path / "nope.xlsx")

    def test_non_xlsx_raises_file_error(self, tmp_path):
        path = tmp_path / "products.csv"
        path.write_text("Product Code,Product Name\nA,B", encoding="utf-8")
        with pytest.raises(FileError):
            import_service.preview_import(path)

    def test_workbook_without_a_header_row_raises_file_error(self, tmp_path):
        path = write_workbook(tmp_path / "bad.xlsx", [GOOD_ROW], header=("Code", "Item"))
        with pytest.raises(FileError, match="header row"):
            import_service.preview_import(path)

    def test_missing_required_column_is_reported(self, tmp_path):
        path = write_workbook(
            tmp_path / "partial.xlsx",
            [("DET001", "Ariel")],
            header=("Product Code", "Product Name"),
        )
        with pytest.raises(FileError, match="missing required columns"):
            import_service.preview_import(path)

    def test_blank_rows_are_skipped(self, tmp_path):
        blank = ("",) * len(TEMPLATE_COLUMNS)
        path = write_workbook(tmp_path / "blanks.xlsx", [GOOD_ROW, blank])
        preview = import_service.preview_import(path)
        assert preview.total_rows == 1


class TestValidation:
    def test_valid_row_is_classified_as_add(self, tmp_path):
        preview = import_service.preview_import(write_workbook(tmp_path / "ok.xlsx", [GOOD_ROW]))
        assert len(preview.add) == 1
        assert not preview.update and not preview.rejected
        assert preview.categories_to_create == ["Detergents"]
        row = preview.add[0]
        assert row.cost_price == Decimal("290.00")
        assert row.opening_stock == 40
        assert row.vat_applicable is True
        assert (row.subcategory, row.brand) == ("Washing Powder", "Ariel")

    def test_category_and_unit_are_normalized(self, tmp_path):
        row = list(GOOD_ROW)
        row[2] = "  detergents "
        row[3] = "  washing   powder "
        row[4] = " ariel "
        row[5] = "PIECE"
        preview = import_service.preview_import(write_workbook(tmp_path / "n.xlsx", [tuple(row)]))
        assert preview.add[0].category == "detergents"
        assert preview.add[0].subcategory == "washing powder"
        assert preview.add[0].brand == "ariel"
        assert preview.add[0].unit == "Piece"

    def test_subcategory_and_brand_may_be_blank(self, tmp_path):
        row = list(GOOD_ROW)
        row[3] = ""
        row[4] = ""
        preview = import_service.preview_import(write_workbook(tmp_path / "blank-levels.xlsx", [tuple(row)]))
        assert not preview.rejected
        assert (preview.add[0].subcategory, preview.add[0].brand) == ("", "")

        import_service.commit_import(preview)
        product = product_service.find_by_code("DET001")
        assert product.subcategory is None
        assert product.brand is None

    @pytest.mark.parametrize(
        "column_index,value,column_name",
        [
            (6, "free", "Cost Price"),
            (7, -20, "Selling Price"),
            (8, 4.5, "Opening Stock"),
            (9, -3, "Reorder Level"),
            (10, "maybe", "VAT Applicable"),
        ],
    )
    def test_bad_cells_are_reported_by_row_and_column(
        self, tmp_path, column_index, value, column_name
    ):
        row = list(GOOD_ROW)
        row[column_index] = value
        preview = import_service.preview_import(write_workbook(tmp_path / "bad.xlsx", [tuple(row)]))

        assert preview.rejected and not preview.add
        error = preview.rejected[0].errors[0]
        assert error.column == column_name
        assert error.row == 2
        assert not preview.is_valid

    def test_overlong_hierarchy_values_are_rejected_by_name(self, tmp_path):
        row = list(GOOD_ROW)
        row[3] = "x" * 121
        row[4] = "y" * 121
        preview = import_service.preview_import(write_workbook(tmp_path / "long.xlsx", [tuple(row)]))
        assert {error.column for error in preview.rejected[0].errors} == {"Subcategory", "Brand"}

    def test_required_blank_cells_are_reported(self, tmp_path):
        row = list(GOOD_ROW)
        row[0] = ""
        row[2] = ""
        row[5] = ""
        row[10] = ""
        preview = import_service.preview_import(write_workbook(tmp_path / "blank.xlsx", [tuple(row)]))
        columns = {error.column for error in preview.rejected[0].errors}
        assert {"Product Code", "Category", "Unit", "VAT Applicable"} <= columns

    def test_all_errors_in_a_row_are_collected(self, tmp_path):
        row = ["", "", "", "", "", "", "free", -1, "x", -2, "maybe"]
        preview = import_service.preview_import(write_workbook(tmp_path / "many.xlsx", [row]))
        assert len(preview.rejected[0].errors) >= 6

    def test_duplicate_code_within_the_workbook_is_rejected(self, tmp_path):
        second = list(GOOD_ROW)
        second[1] = "2kg Washing Powder"
        preview = import_service.preview_import(
            write_workbook(tmp_path / "dupes.xlsx", [GOOD_ROW, tuple(second)])
        )
        assert len(preview.add) == 1
        assert len(preview.rejected) == 1
        assert "Duplicate product code" in preview.rejected[0].errors[0].message

    def test_error_report_can_be_exported(self, tmp_path):
        preview = import_service.preview_import(
            write_workbook(
                tmp_path / "bad.xlsx",
                [("X", "No prices", "Cat", "", "", "Piece", "a", "b", 1, 1, "Yes")],
            )
        )
        report = import_service.write_error_report(preview, tmp_path / "errors.xlsx")
        sheet = load_workbook(report)["Import Errors"]
        assert sheet.cell(row=1, column=3).value == "Problem"
        assert sheet.max_row >= 2


class TestCommit:
    def test_import_creates_products_categories_and_opening_movements(self, tmp_path):
        rows = [
            GOOD_ROW,
            ("BRD001", "Supreme Bread 400g", "Bakery", "Bread", "Supreme", "Piece", 45, 60, 25, 8, "No"),
            ("WTR001", "Mineral Water 1L", "Beverages", "Water", "", "Piece", 30, 50, 0, 5, "Yes"),
        ]
        summary = import_service.commit_import(
            import_service.preview_import(write_workbook(tmp_path / "ok.xlsx", rows))
        )

        assert summary.imported == 3
        assert summary.errors == 0
        assert sorted(summary.created_categories) == ["Bakery", "Beverages", "Detergents"]
        assert product_service.catalogue_stats() == {
            "products": 3,
            "categories": 3,
            "archived": 0,
        }

        ariel = product_service.find_by_code("DET001")
        assert ariel.stock_quantity == 40
        assert ariel.vat_applicable is True
        assert (ariel.subcategory, ariel.brand) == ("Washing Powder", "Ariel")

        # Opening stock is traceable in the ledger; the zero-stock product has none.
        movements = inventory_service.movements_for_product(ariel.id)
        assert [(m.reason, m.quantity, m.direction) for m in movements] == [
            ("Initial Stock", 40, "in")
        ]
        assert inventory_service.movements_for_product(
            product_service.find_by_code("WTR001").id
        ) == []

    def test_nothing_is_written_when_the_preview_has_errors(self, tmp_path):
        preview = import_service.preview_import(
            write_workbook(
                tmp_path / "mixed.xlsx",
                [GOOD_ROW, ("B1", "Bad", "Cat", "", "", "Piece", "x", 1, 1, 1, "Yes")],
            )
        )
        assert preview.rejected
        # The UI stops the user here; committing an invalid preview must not create rows.
        assert not preview.is_valid
        assert product_service.catalogue_stats()["products"] == 0

    def test_reimport_updates_master_data_but_never_stock(self, tmp_path):
        import_service.commit_import(
            import_service.preview_import(write_workbook(tmp_path / "first.xlsx", [GOOD_ROW]))
        )
        product = product_service.find_by_code("DET001")
        inventory_service.record_movement(
            product_id=product.id,
            movement_type="Stock Out",
            reason="Damaged",
            quantity=15,
            notes="Crushed",
        )
        product_service.set_image(product.id, "products/keep-me.jpg")
        stock_before = product_service.find_by_code("DET001").stock_quantity
        assert stock_before == 25

        changed = (
            "DET001", "2kg Washing Powder", "Laundry", "Powder", "Omo",
            "Pack", 310, 399, 999, 12, "No",
        )
        preview = import_service.preview_import(write_workbook(tmp_path / "second.xlsx", [changed]))
        assert len(preview.update) == 1 and not preview.add
        assert any(change.startswith("Selling Price") for change in preview.update[0].changes)
        assert "Brand: Ariel -> Omo" in preview.update[0].changes

        summary = import_service.commit_import(preview)
        assert summary.updated == 1 and summary.imported == 0

        after = product_service.find_by_code("DET001")
        assert after.name == "2kg Washing Powder"
        assert after.category_name == "Laundry"
        assert (after.subcategory, after.brand) == ("Powder", "Omo")
        assert after.unit == "Pack"
        assert after.cost_price == Decimal("310.00")
        assert after.selling_price == Decimal("399.00")
        assert after.reorder_level == 12
        assert after.vat_applicable is False
        # The critical rule: live stock and image survived the import untouched.
        assert after.stock_quantity == stock_before
        assert after.image_path == "products/keep-me.jpg"
        assert product_service.catalogue_stats()["products"] == 1

    def test_existing_category_is_reused_not_duplicated(self, tmp_path):
        import_service.commit_import(
            import_service.preview_import(write_workbook(tmp_path / "a.xlsx", [GOOD_ROW]))
        )
        second = ("DET002", "2kg Washing Powder", "detergents", "Washing Powder", "Omo",
                  "Piece", 500, 600, 5, 2, "Yes")
        summary = import_service.commit_import(
            import_service.preview_import(write_workbook(tmp_path / "b.xlsx", [second]))
        )
        assert summary.created_categories == []
        assert product_service.catalogue_stats()["categories"] == 1

    def test_commit_is_atomic_when_a_row_fails_mid_transaction(self, tmp_path, monkeypatch):
        rows = [
            GOOD_ROW,
            ("BRD001", "Supreme Bread 400g", "Bakery", "Bread", "Supreme", "Piece", 45, 60, 25, 8, "No"),
        ]
        preview = import_service.preview_import(write_workbook(tmp_path / "atomic.xlsx", rows))

        original = import_service.apply_movement
        calls = {"count": 0}

        def explode(*args, **kwargs):
            calls["count"] += 1
            if calls["count"] == 2:
                raise RuntimeError("simulated disk failure")
            return original(*args, **kwargs)

        monkeypatch.setattr(import_service, "apply_movement", explode)

        from core.exceptions import DatabaseError

        with pytest.raises(DatabaseError):
            import_service.commit_import(preview)

        assert product_service.catalogue_stats()["products"] == 0
        assert product_service.list_categories() == []
        assert inventory_service.list_movements() == []


class TestHierarchy:
    """Category -> Subcategory -> Brand through the import pipeline."""

    COFFEE_ROWS = (
        ("COF001", "Classic", "Beverages", "Coffee", "Nescafé", "Piece", 480, 570, 24, 6, "Yes"),
        ("COF002", "Gold", "Beverages", "Coffee", "Nescafé", "Piece", 850, 990, 12, 4, "Yes"),
        ("COF003", "House Blend", "Beverages", "Coffee", "Dormans", "Pack", 620, 750, 18, 5, "Yes"),
    )

    def test_brands_group_under_their_subcategory(self, tmp_path):
        summary = import_service.commit_import(
            import_service.preview_import(
                write_workbook(tmp_path / "coffee.xlsx", list(self.COFFEE_ROWS))
            )
        )
        assert (summary.imported, summary.errors) == (3, 0)

        assert product_service.subcategory_names() == ["Coffee"]
        assert product_service.brand_names() == ["Dormans", "Nescafé"]
        assert product_service.brand_names(subcategory="Coffee") == ["Dormans", "Nescafé"]

        hierarchy = product_service.catalogue_hierarchy()
        coffee = hierarchy["Beverages"]["Coffee"]
        assert sorted(coffee) == ["Dormans", "Nescafé"]
        assert [product.code for product in coffee["Nescafé"]] == ["COF001", "COF002"]
        assert [product.code for product in coffee["Dormans"]] == ["COF003"]

    def test_legacy_workbook_imports_with_the_levels_blank(self, tmp_path):
        assert LEGACY_HEADER == REQUIRED_COLUMNS
        preview = import_service.preview_import(
            write_workbook(tmp_path / "old.xlsx", [LEGACY_ROW], header=LEGACY_HEADER)
        )
        assert preview.missing_optional == ["Subcategory", "Brand"]
        assert (preview.add[0].subcategory, preview.add[0].brand) == (None, None)

        summary = import_service.commit_import(preview)
        assert summary.imported == 1
        product = product_service.find_by_code("DET001")
        assert product.name == "Ariel Detergent 1kg"
        assert product.subcategory is None
        assert product.brand is None

    def test_legacy_workbook_never_wipes_levels_it_does_not_carry(self, tmp_path):
        import_service.commit_import(
            import_service.preview_import(write_workbook(tmp_path / "new.xlsx", [GOOD_ROW]))
        )
        preview = import_service.preview_import(
            write_workbook(tmp_path / "old.xlsx", [LEGACY_ROW], header=LEGACY_HEADER)
        )
        assert len(preview.update) == 1
        assert not any(change.startswith(("Subcategory", "Brand")) for change in preview.update[0].changes)

        import_service.commit_import(preview)
        product = product_service.find_by_code("DET001")
        assert product.name == "Ariel Detergent 1kg"
        assert (product.subcategory, product.brand) == ("Washing Powder", "Ariel")

    def test_columns_are_matched_by_header_name_not_position(self, tmp_path):
        shuffled = (
            "Brand", "Product Code", "VAT Applicable", "Category", "Selling Price",
            "Subcategory", "Unit", "Product Name", "Cost Price", "Reorder Level", "Opening Stock",
        )
        row = (
            "Nescafé", "COF001", "Yes", "Beverages", 570,
            "Coffee", "Piece", "Classic", 480, 6, 24,
        )
        preview = import_service.preview_import(
            write_workbook(tmp_path / "shuffled.xlsx", [row], header=shuffled)
        )
        assert not preview.rejected
        assert preview.add[0].code == "COF001"
        assert preview.add[0].name == "Classic"
        assert (preview.add[0].subcategory, preview.add[0].brand) == ("Coffee", "Nescafé")
        assert preview.add[0].selling_price == Decimal("570.00")
        assert preview.add[0].opening_stock == 24
